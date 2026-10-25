"""Send the Friday email from GitHub Actions.

Exports the tracker sheet as .xlsx with the service account, builds the email with
weekly_email.build() and the printable PDF with print_sheet.build_pdf(), checks the
email, and sends both through Gmail SMTP.

Environment:
  GOOGLE_SERVICE_ACCOUNT_JSON  the same key the tracker uses
  GMAIL_APP_PASSWORD           a Google app password for EMAIL_FROM
  EMAIL_FROM / EMAIL_TO        default zennode@gmail.com

  python send_weekly_email.py [--week-of 2026-10-26] [--dry-run out.html]
"""

import argparse
import html
import io
import json
import os
import re
import smtplib
import sys
from datetime import date, timedelta
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import openpyxl
import yaml
from google.auth.transport.requests import AuthorizedSession
from google.oauth2.service_account import Credentials

import print_sheet
import weekly_email

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GRADIENT = "linear-gradient(135deg,#1e3a5f,#2d5a8e)"


class _Tab:
    def __init__(self, values):
        self.values = values

    def iter_rows(self, values_only=True):
        return iter(self.values)


class _Book:
    """Just enough of an openpyxl workbook for weekly_email.read_tab()."""

    def __init__(self, tabs):
        self.tabs = tabs
        self.sheetnames = list(tabs)

    def __getitem__(self, name):
        return _Tab(self.tabs[name])


def download_sheet(spreadsheet_id):
    creds = Credentials.from_service_account_info(
        json.loads(os.environ["GOOGLE_SERVICE_ACCOUNT_JSON"]),
        scopes=["https://www.googleapis.com/auth/drive.readonly",
                "https://www.googleapis.com/auth/spreadsheets.readonly"])
    resp = AuthorizedSession(creds).get(
        f"https://www.googleapis.com/drive/v3/files/{spreadsheet_id}/export",
        params={"mimeType": XLSX}, timeout=120)
    if resp.ok:
        return openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)
    # The Drive API may not be enabled on the service account's project; the Sheets
    # API is (the tracker uses it), so read the tabs as displayed text instead.
    print(f"Drive export failed ({resp.status_code}); reading tabs through the Sheets API")
    import gspread
    book = gspread.authorize(creds).open_by_key(spreadsheet_id)
    names = ["Auto - Auction Tracker", "Auto - Cancelled & Postponed", "Auto - Status Changes"]
    return _Book({ws.title: ws.get_all_values() for ws in book.worksheets() if ws.title in names})


def check(body, summary):
    """The payload checks from the email-format standard; returns a list of failures."""
    plain = re.sub(r"\s+", " ", re.sub(r"<[^>]*>", " ", body)).strip()
    problems = []
    if body.count(weekly_email.TITLE) != 1:
        problems.append(f"title appears {body.count(weekly_email.TITLE)} times")
    if GRADIENT not in body:
        problems.append("header gradient missing")
    if summary["auctions"]:
        if "<table" not in body:
            problems.append("no table")
        rows = len(re.findall(r"<tr[\s>]", body))
        if rows < summary["min_rows"]:
            problems.append(f"{rows} table rows, expected {summary['min_rows']}")
    if len(plain) < 300:
        problems.append(f"only {len(plain)} characters of text")
    if re.search(r"Matthew Tortoriello|It's just you against you", body):
        problems.append("signature text in body")
    return problems


def send(subject, body, sender, to, password, pdf=None, pdf_name="auctions.pdf"):
    msg = MIMEMultipart("mixed")
    msg["Subject"], msg["From"], msg["To"] = subject, sender, to
    alt = MIMEMultipart("alternative")
    text = html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]*>", " ", body)).strip())
    alt.attach(MIMEText(text, "plain"))
    alt.attach(MIMEText(body, "html"))
    msg.attach(alt)
    if pdf:
        part = MIMEApplication(pdf, _subtype="pdf")
        part.add_header("Content-Disposition", "attachment", filename=pdf_name)
        msg.attach(part)
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=60) as s:
        s.login(sender, password)
        s.sendmail(sender, [to], msg.as_string())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--week-of", help="build for the week starting this Monday (YYYY-MM-DD)")
    ap.add_argument("--dry-run", metavar="OUT", help="write the HTML here instead of sending")
    args = ap.parse_args()

    with open(os.path.join(os.path.dirname(__file__), "config.yaml")) as f:
        spreadsheet_id = yaml.safe_load(f)["spreadsheet_id"]
    today = date.today()
    if args.week_of:
        today = date.fromisoformat(args.week_of) - timedelta(days=3)

    wb = download_sheet(spreadsheet_id)
    body, summary = weekly_email.build(wb, today)
    print(json.dumps(summary, indent=1))
    start, end, week = weekly_email.select_week(wb, today)
    pdf = print_sheet.build_pdf(week, start, end)
    pdf_name = f"Auctions {start:%b %-d} - {end:%b %-d}.pdf"
    problems = check(body, summary)
    if problems:
        sys.exit("Not sending: " + "; ".join(problems))

    if args.dry_run:
        with open(args.dry_run, "w") as f:
            f.write(body)
        with open(os.path.splitext(args.dry_run)[0] + ".pdf", "wb") as f:
            f.write(pdf)
        return
    sender = os.environ.get("EMAIL_FROM") or "zennode@gmail.com"
    to = os.environ.get("EMAIL_TO") or "zennode@gmail.com"
    password = os.environ.get("GMAIL_APP_PASSWORD")
    if not password:
        sys.exit("Set the GMAIL_APP_PASSWORD secret to send the email")
    send(summary["subject"], body, sender, to, password.replace(" ", ""), pdf, pdf_name)
    print(f"Sent \"{summary['subject']}\" to {to} with {pdf_name} ({len(pdf):,} bytes)")


if __name__ == "__main__":
    main()
